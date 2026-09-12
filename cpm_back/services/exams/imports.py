"""Bounded XLSX previews and atomic, creator-scoped exam imports."""

from collections import Counter
from datetime import datetime, date, timedelta
from io import BytesIO
import hashlib
import re
from zipfile import ZipFile, BadZipFile

from . import admin
from .common import (
    ExamError,
    fail,
    fields,
    integer,
    text_value,
    decimal_number,
    dumps,
    loads,
    now,
    iso,
    digest,
    check_version,
    page_args,
    pagination,
    flush_invalidations,
)

TABLES = {
    "questions": "classic_exam_question_import_sessions",
    "assignments": "classic_exam_assignment_import_sessions",
    "outside": "outside_exam_result_import_sessions",
}
MAX_FILE = 10 * 1024 * 1024
MAX_PREVIEW = 16 * 1024 * 1024


def canonical_header(value, kind):
    if not isinstance(value, str):
        fail("invalid_import_headers", "Заголовки должны быть текстом")
    value = value.strip().casefold()
    common = {"student_id": "studentId", "student_login": "studentLogin"}
    if kind == "questions":
        aliases = {
            "часть": "partCode",
            "part_code": "partCode",
            "вопрос": "questionText",
            "question_text": "questionText",
            "эталонный ответ": "answerText",
            "answer_text": "answerText",
        }
    elif kind == "outside":
        aliases = dict(common, points="points", grade="grade", examinator="examinator")
    else:
        aliases = dict(common, replacement_limit="replacementLimit")
        for n in range(1, 7):
            aliases[f"examinator_{n}_id"] = f"examinator{n}Id"
            for key in (f"examinator_{n}", f"examinator_{n}_login", f"экзаменатор {n}"):
                aliases[key] = f"examinator{n}Login"
    if value not in aliases:
        fail("unknown_import_header", "Неизвестный заголовок", header=value)
    return aliases[value]


def cell_value(cell):
    value = cell.value
    if value is None:
        return None
    if cell.data_type in ("f", "e") or isinstance(value, (bool, date, datetime)):
        return {"invalidExcelCell": cell.data_type, "value": str(value)}
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value


def parse_xlsx(content, filename, kind):
    if kind not in TABLES:
        fail("invalid_import_type")
    if not filename or not filename.casefold().endswith(".xlsx"):
        fail("invalid_import_file", "Поддерживается только .xlsx")
    if len(content) > MAX_FILE:
        fail("payload_too_large", "Файл должен быть не больше10MiB", 413)
    try:
        with ZipFile(BytesIO(content)) as archive:
            members = archive.infolist()
            if (
                len(members) > 10000
                or sum(m.file_size for m in members) > 64 * 1024 * 1024
            ):
                fail("payload_too_large", "Превышен размер распакованного Excel", 413)
            if any(m.flag_bits & 1 for m in members):
                fail("invalid_import_file", "Зашифрованный Excel не поддерживается")
            # Disallow DTD/entity declarations before passing XML to openpyxl.
            for member in members:
                if not member.filename.endswith(".xml"):
                    continue
                with archive.open(member) as stream:
                    tail = b""
                    while True:
                        block = stream.read(65536)
                        if not block:
                            break
                        scan = (tail + block).replace(b"\x00", b"").upper()
                        if b"<!DOCTYPE" in scan or b"<!ENTITY" in scan:
                            fail(
                                "invalid_import_file",
                                "DTD и XML entities не поддерживаются",
                            )
                        tail = block[-64:]
        from openpyxl import load_workbook

        workbook = load_workbook(
            BytesIO(content), read_only=True, data_only=False, keep_links=False
        )
    except ExamError:
        raise
    except Exception:
        fail("invalid_import_file", "Не удалось прочитать Excel")
    try:
        sheets = []
        for sheet in workbook.worksheets:
            if sheet.max_column and sheet.max_column > 32:
                fail("import_columns_limit", "Разрешено не больше32колонок")
            data_rows = []
            scanned = 0
            for source_row, cells in enumerate(sheet.iter_rows(), 1):
                scanned += 1
                if scanned > 100000:
                    fail("import_rows_limit", "Слишком много строк Excel")
                values = [cell_value(c) for c in cells]
                if not any(v is not None and v != "" for v in values):
                    continue
                data_rows.append((source_row, values))
                if len(data_rows) > 5001:
                    fail("import_rows_limit", "Разрешено не больше5000строк данных")
            if data_rows:
                sheets.append(data_rows)
            if len(sheets) > 1:
                fail("multiple_import_sheets", "Оставьте один непустой лист")
        if not sheets or len(sheets[0]) < 2:
            fail("import_no_rows", "В файле нет строк данных", 422)
        data = sheets[0]
        if data[0][0] != 1:
            fail(
                "invalid_import_headers", "Заголовки должны находиться в первой строке"
            )
        raw_headers = data[0][1]
        while raw_headers and raw_headers[-1] is None:
            raw_headers.pop()
        headers = [canonical_header(value, kind) for value in raw_headers]
        if len(set(headers)) != len(headers):
            fail("duplicate_import_headers", "Заголовки дублируют одно поле")
        required = (
            {"partCode", "questionText", "answerText"}
            if kind == "questions"
            else {"points", "grade", "examinator"} if kind == "outside" else set()
        )
        if not required.issubset(headers) or (
            kind != "questions"
            and not {"studentId", "studentLogin"}.intersection(headers)
        ):
            fail("missing_import_headers", "Отсутствуют обязательные колонки")
        result = []
        for source_row, values in data[1:]:
            if any(
                value is not None and value != "" for value in values[len(headers) :]
            ):
                fail(
                    "invalid_import_headers", "Данные находятся в колонке без заголовка"
                )
            input = {
                header: value
                for header, value in zip(headers, values)
                if value is not None and value != ""
            }
            result.append(
                {
                    "rowId": f"row-{source_row}",
                    "sourceRow": source_row,
                    "excluded": False,
                    "input": input,
                }
            )
        return result
    finally:
        workbook.close()


def resolve_person(db, role, id_value=None, login=None, require_login=True, cache=None):
    identifier = None
    if id_value not in (None, ""):
        if isinstance(id_value, str) and re.fullmatch(r"[0-9]+", id_value):
            id_value = int(id_value)
        identifier = integer(id_value, role + "_id")
    if login not in (None, ""):
        # An explicit numeric login is still a login, never a role entity ID.
        if type(login) in (int, float):
            login = str(login)
        login = text_value(login, role + "_login", 50)
        lookup_key = (role, login)
        if cache is not None and lookup_key in cache["logins"]:
            found = cache["logins"][lookup_key]
        else:
            found = db.one(
                "SELECT ref_id FROM auth_users WHERE role=%s AND username=%s",
                (role, login),
            )
            if cache is not None:
                cache["logins"][lookup_key] = found
        if not found:
            fail(role + "_not_found", "Аккаунт не найден")
        if identifier is not None and identifier != found["ref_id"]:
            fail(
                "identifier_login_conflict",
                "ID и логин относятся к разным пользователям",
            )
        identifier = found["ref_id"]
    if identifier is None:
        fail(role + "_required", "Укажите ID или логин")
    if cache is None:
        return admin.account(db, role, identifier, require_login)
    if role not in cache:
        table = {"student": "students", "examinator": "examinators"}[role]
        cache[role] = {
            p["id"]: p
            for p in db.all(
                f"""SELECT u.id,u.full_name,
            (SELECT MIN(username) FROM auth_users a WHERE a.ref_id=u.id AND a.role=%s) login
            FROM {table} u""",
                (role,),
            )
        }
    person = cache[role].get(identifier)
    if not person or (require_login and person["login"] is None):
        fail(role + "_not_found", "Пользователь не найден или не имеет аккаунта")
    return person


def _issue(error):
    return {
        "code": error.code,
        "field": error.details.get("field"),
        "message": error.message,
        "details": error.details,
    }


def validate_rows(db, exam_id, kind, rows, refresh_baselines=False):
    parts = (
        {p["code"]: p for p in admin.part_rows(db, exam_id)}
        if kind == "questions"
        else {}
    )
    cache = {"logins": {}}
    existing_questions = (
        {
            (q["part_id"], q["content_hash"])
            for q in db.all(
                "SELECT part_id,content_hash FROM classic_exam_questions WHERE exam_id=%s",
                (exam_id,),
            )
        }
        if kind == "questions"
        else set()
    )
    existing_students = (
        {
            p["student_id"]
            for p in db.all(
                "SELECT student_id FROM "
                + ("exam_sessions" if kind == "outside" else "classic_exam_assignments")
                + " WHERE exam_id=%s"
                + (" AND attempt_no=1" if kind == "assignments" else ""),
                (exam_id,),
            )
        }
        if kind != "questions"
        else set()
    )
    privileges = (
        {
            p["student_id"]: p
            for p in db.all(
                "SELECT student_id,replacement_limit,version FROM classic_exam_student_privileges WHERE exam_id=%s",
                (exam_id,),
            )
        }
        if kind == "assignments"
        else {}
    )
    fingerprints = {}
    validated = []
    for raw in rows:
        row = dict(
            raw,
            resolved={},
            errors=[],
            warnings=[],
            action="skip" if raw["excluded"] else "create",
        )
        if row["excluded"]:
            validated.append(row)
            continue
        try:
            values = row["input"]
            if kind == "questions":
                fields(
                    values,
                    ("partCode", "questionText", "answerText"),
                    ("partCode", "questionText", "answerText"),
                )
                code = text_value(values["partCode"], "partCode", 1).upper()
                if not re.fullmatch("[A-Z]", code):
                    fail("invalid_part_code", "Часть — латинская буква A–Z")
                question = text_value(
                    values["questionText"], "questionText", None, 61440
                )
                answer = text_value(values["answerText"], "answerText", None, 61440)
                row["resolved"] = {
                    "partCode": code,
                    "partId": parts.get(code, {}).get("id"),
                    "questionText": question,
                    "answerText": answer,
                }
                key = (code, digest([question, answer]))
                if code in parts and (parts[code]["id"], key[1]) in existing_questions:
                    fail("question_duplicate", "Вопрос уже есть в этой части")
            else:
                allowed = (
                    ("studentId", "studentLogin", "points", "grade", "examinator")
                    if kind == "outside"
                    else ("studentId", "studentLogin", "replacementLimit")
                    + tuple(
                        f"examinator{n}{suffix}"
                        for n in range(1, 7)
                        for suffix in ("Id", "Login")
                    )
                )
                fields(values, allowed)
                student = resolve_person(
                    db,
                    "student",
                    values.get("studentId"),
                    values.get("studentLogin"),
                    kind != "outside",
                    cache,
                )
                row["resolved"]["student"] = {
                    "id": student["id"],
                    "fullName": student["full_name"],
                    "login": student["login"],
                }
                key = student["id"]
                if kind == "outside":
                    row["resolved"].update(
                        points=decimal_number(values.get("points")),
                        grade=integer(values.get("grade"), "grade", 0, 5),
                        examinator=text_value(
                            values.get("examinator"), "examinator", 255
                        ),
                    )
                    if key in existing_students:
                        fail(
                            "outside_result_already_exists",
                            "У студента уже есть результат",
                        )
                else:
                    members = []
                    for n in range(1, 7):
                        member_id, login = values.get(f"examinator{n}Id"), values.get(
                            f"examinator{n}Login"
                        )
                        if member_id in (None, "") and login in (None, ""):
                            continue
                        person = resolve_person(
                            db, "examinator", member_id, login, cache=cache
                        )
                        members.append(
                            {
                                "id": person["id"],
                                "fullName": person["full_name"],
                                "login": person["login"],
                            }
                        )
                    if not 1 <= len(members) <= 6 or len(
                        {m["id"] for m in members}
                    ) != len(members):
                        fail(
                            "invalid_commission_size",
                            "Нужны от1до6 разных экзаменаторов",
                        )
                    row["resolved"]["members"] = members
                    if key in existing_students:
                        fail("assignment_already_exists", "Студент уже назначен")
                    if "replacementLimit" in values:
                        row["resolved"]["replacementLimit"] = integer(
                            values["replacementLimit"], "replacementLimit", 0, 100
                        )
                        current = privileges.get(key)
                        privilege = {
                            "version": current["version"] if current else 0,
                            "replacementLimit": (
                                current["replacement_limit"] if current else 0
                            ),
                        }
                        refresh = refresh_baselines is True or (
                            isinstance(refresh_baselines, set)
                            and row["rowId"] in refresh_baselines
                        )
                        if refresh or "privilegeBaselineVersion" not in row:
                            row["privilegeBaselineVersion"] = privilege["version"]
                            row["privilegeBaselineStudentId"] = key
                        elif (
                            row["privilegeBaselineVersion"] != privilege["version"]
                            or row.get("privilegeBaselineStudentId") != key
                        ):
                            fail(
                                "privilege_modified",
                                "Привилегия изменилась; пересохраните превью",
                            )
                        row["warnings"].append(
                            {
                                "code": "privilege_change",
                                "field": "replacementLimit",
                                "message": "Лимит замен будет установлен",
                                "details": {
                                    "oldValue": privilege["replacementLimit"],
                                    "newValue": values["replacementLimit"],
                                },
                            }
                        )
            if key in fingerprints:
                issue = {
                    "code": "duplicate_import_row",
                    "field": None,
                    "message": "Повторная строка в файле",
                    "details": {},
                }
                row["errors"].append(issue)
                fingerprints[key]["errors"].append(issue.copy())
            else:
                fingerprints[key] = row
        except ExamError as error:
            row["errors"].append(_issue(error))
        validated.append(row)
    return validated


def summary(rows):
    included = [r for r in rows if not r["excluded"]]
    errors = sum(bool(r["errors"]) for r in included)
    return {
        "total": len(rows),
        "included": len(included),
        "excluded": len(rows) - len(included),
        "creatable": len(included) - errors,
        "conflicts": errors,
        "errors": errors,
    }


def stored_rows(rows):
    keys = (
        "rowId",
        "sourceRow",
        "excluded",
        "input",
        "privilegeBaselineVersion",
        "privilegeBaselineStudentId",
    )
    return [{key: row[key] for key in keys if key in row} for row in rows]


def payload_json(db, rows):
    value = dumps(stored_rows(rows))
    # Fail clearly before the driver/server connection is killed by packet limits.
    packet = db.scalar("SELECT @@max_allowed_packet")
    escaped_bytes = (
        len(value.encode("utf-8")) + value.count("\\") + value.count("'") + 4096
    )
    if len(value.encode("utf-8")) > MAX_PREVIEW or escaped_bytes > packet:
        fail(
            "payload_too_large",
            "Превью превышает лимит размера сервера",
            413,
            maxPreviewBytes=MAX_PREVIEW,
            serverPacketBytes=packet,
        )
    return value


def session_row(db, kind, session_id, actor, exam_id=None, lock=False):
    row = db.one(
        f"SELECT * FROM {TABLES[kind]} WHERE id=%s" + (" FOR UPDATE" if lock else ""),
        (session_id,),
    )
    if (
        not row
        or row["created_by_role"] != actor["role"]
        or row["created_by"] != actor["id"]
        or (exam_id is not None and row["exam_id"] != exam_id)
    ):
        fail("import_session_not_found", "Сессия импорта не найдена", 404)
    if row["expires_at"] < now():
        fail("import_session_expired", "Сессия импорта истекла", 410)
    return row


def session_dto(db, kind, row, args=None, changes=None):
    rows = (
        validate_rows(db, row["exam_id"], kind, loads(row["preview_payload"]))
        if row["status"] == "editable"
        else loads(row["preview_payload"])
    )
    result = {
        "id": row["id"],
        "examId": row["exam_id"],
        "previewVersion": row["preview_version"],
        "status": row["status"],
        "sourceFilename": row["source_filename"],
        "expiresAt": iso(row["expires_at"]),
        "commitResult": loads(row["commit_result"]),
    }
    if row["status"] == "committed":
        included = len([r for r in rows if not r["excluded"]])
        result.update(
            rows=[],
            summary={
                "total": len(rows),
                "included": included,
                "excluded": len(rows) - included,
                "creatable": 0,
                "conflicts": 0,
                "errors": 0,
            },
        )
        return result
    result["summary"] = summary(rows)
    if changes is not None:
        result["rows"] = [r for r in rows if r["rowId"] in changes]
    else:
        args = args or {}
        page, limit, offset = page_args(args)
        if args.get("onlyErrors") not in (None, "true", "false"):
            fail("invalid_only_errors")
        visible = (
            [r for r in rows if r["errors"]]
            if args.get("onlyErrors") == "true"
            else rows
        )
        result["rows"] = visible[offset : offset + limit]
        result["pagination"] = pagination(len(visible), page, limit)
    # Internal baseline tokens never become client-writable fields.
    for item in result["rows"]:
        item.pop("privilegeBaselineVersion", None)
        item.pop("privilegeBaselineStudentId", None)
    return result


def create_session(db, kind, exam_id, actor, filename, rows):
    validated = validate_rows(db, exam_id, kind, rows, True)
    session_id = db.insert(
        TABLES[kind],
        {
            "exam_id": exam_id,
            "created_by_role": actor["role"],
            "created_by": actor["id"],
            "source_filename": text_value(filename, "filename", 255),
            "preview_payload": payload_json(db, validated),
            "expires_at": now() + timedelta(hours=72),
        },
    )
    return session_dto(db, kind, session_row(db, kind, session_id, actor, exam_id))


def update_session(db, kind, row, payload):
    fields(
        payload,
        ("expectedPreviewVersion", "changes"),
        ("expectedPreviewVersion", "changes"),
    )
    if row["status"] != "editable":
        fail("import_already_committed", status=409)
    check_version(
        row,
        payload["expectedPreviewVersion"],
        "import_preview_modified",
        "preview_version",
    )
    changes = payload["changes"]
    if not isinstance(changes, list) or len(changes) > 5000:
        fail("invalid_preview_changes")
    rows = loads(row["preview_payload"])
    by_id = {r["rowId"]: r for r in rows}
    edited = set()
    for change in changes:
        fields(change, ("rowId", "input", "excluded"), ("rowId",))
        key = change["rowId"]
        if not isinstance(key, str) or key not in by_id or key in edited:
            fail("invalid_preview_row")
        edited.add(key)
        if "input" in change:
            if not isinstance(change["input"], dict):
                fail("invalid_preview_input")
            by_id[key]["input"] = change["input"]
        if "excluded" in change:
            if type(change["excluded"]) is not bool:
                fail("invalid_preview_excluded")
            by_id[key]["excluded"] = change["excluded"]
    validated = validate_rows(db, row["exam_id"], kind, rows, edited)
    encoded = payload_json(db, validated)
    version = row["preview_version"] + (encoded != dumps(loads(row["preview_payload"])))
    db.update(
        TABLES[kind],
        {"preview_payload": encoded, "preview_version": version},
        "id=%s",
        (row["id"],),
    )
    row = dict(row, preview_payload=encoded, preview_version=version)
    return session_dto(db, kind, row, changes=edited)


def commit_session(db, kind, row, actor, payload):
    fields(payload, ("expectedPreviewVersion",), ("expectedPreviewVersion",))
    integer(payload["expectedPreviewVersion"], "previewVersion")
    if row["status"] == "committed":
        return dict(loads(row["commit_result"]), replayed=True)
    check_version(
        row,
        payload["expectedPreviewVersion"],
        "import_preview_modified",
        "preview_version",
    )
    validated = validate_rows(db, row["exam_id"], kind, loads(row["preview_payload"]))
    included = [r for r in validated if not r["excluded"]]
    conflicts = [r["rowId"] for r in included if r["errors"]]
    if conflicts:
        fail(
            "import_conflict",
            "Исправьте или исключите конфликтные строки",
            409,
            rowIds=conflicts,
        )
    if not included:
        fail("import_no_rows", "Нет строк для импорта", 422)
    counts = {
        key: 0
        for key in (
            "rows",
            "partsCreated",
            "questionsCreated",
            "commissionsCreated",
            "assignmentsCreated",
            "resultsCreated",
            "privilegesChanged",
        )
    }
    counts["rows"] = len(included)
    db.defer_invalidations = True
    exam_id = row["exam_id"]
    parts = (
        {p["code"]: p["id"] for p in admin.part_rows(db, exam_id)}
        if kind == "questions"
        else {}
    )
    commissions = {}
    if kind == "assignments":
        for commission in db.all(
            "SELECT id FROM classic_exam_commissions WHERE exam_id=%s ORDER BY id",
            (exam_id,),
        ):
            members = tuple(
                sorted(m["id"] for m in admin.commission_members(db, commission["id"]))
            )
            commissions.setdefault(members, commission["id"])
    for item in included:
        values = item["resolved"]
        if kind == "questions":
            code = values["partCode"]
            if code not in parts:
                parts[code] = admin.save_part(db, exam_id, {"code": code})["id"]
                counts["partsCreated"] += 1
            admin.save_question(
                db,
                exam_id,
                {
                    "partId": parts[code],
                    "questionText": values["questionText"],
                    "answerText": values["answerText"],
                },
            )
            counts["questionsCreated"] += 1
        elif kind == "outside":
            admin.save_outside_result(
                db,
                exam_id,
                {
                    "studentId": values["student"]["id"],
                    "points": values["points"],
                    "grade": values["grade"],
                    "examinator": values["examinator"],
                },
            )
            counts["resultsCreated"] += 1
        else:
            members = tuple(sorted(m["id"] for m in values["members"]))
            if members not in commissions:
                commissions[members] = admin.save_commission(
                    db, exam_id, {"examinatorIds": list(members)}
                )["id"]
                counts["commissionsCreated"] += 1
            data = {
                "studentId": values["student"]["id"],
                "commissionId": commissions[members],
            }
            if "replacementLimit" in values:
                data.update(
                    replacementLimit=values["replacementLimit"],
                    expectedPrivilegeVersion=item["privilegeBaselineVersion"],
                )
                previous = admin.get_privilege(db, exam_id, data["studentId"])
                counts["privilegesChanged"] += int(
                    previous["version"] == 0
                    or previous["replacementLimit"] != values["replacementLimit"]
                )
            admin.save_assignment(db, exam_id, data, actor)
            counts["assignmentsCreated"] += 1
    result = {
        "sessionId": row["id"],
        "committed": True,
        "counts": counts,
        "replayed": False,
    }
    db.update(
        TABLES[kind],
        {
            "status": "committed",
            "committed_at": now(),
            "expires_at": now() + timedelta(days=7),
            "commit_result": dumps(result),
        },
        "id=%s",
        (row["id"],),
    )
    flush_invalidations(db)
    return result
